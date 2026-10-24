"""Gera a planilha de simulação do IRPF mínimo (Lei 15.270/2025), uma aba por contribuinte.

Uso:
    python tools/planilha_irpfm.py dados.json saida.xlsx

As regras espelham a calculadora web (src/utils/dividends.ts e Calculator.tsx):
- lucros e dividendos em uma tabela de fontes pagadoras, uma linha por CNPJ, com a
  parcela deliberada até 2025 (protegida) separada da posterior (sujeita);
- alocação padrão: tudo como posterior, com ressalva na observação;
- fontes abaixo de R$ 500 agrupadas em "Demais fontes" (sem redutor);
- redutor do art. 16-B fonte a fonte; retenção de 10% opcional por fonte (padrão: Não).

Formato do JSON (lista de contribuintes) — ver README, seção "Planilha de simulação".
Os dados de clientes NÃO devem ser salvos dentro deste repositório.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SMALL_SOURCE_THRESHOLD = 500
DEFAULT_NOMINAL = 0.34
SPARE_SOURCE_ROWS = 3
DEFAULT_ALLOCATION_NOTE = ("Alocado por padrão como posterior a 2025 (sujeito). Se houver deliberação até 31/12/2025, "
                           "informe a parcela protegida na coluna D.")

TRIB_LABELS = [
    "Recebidos de Pessoa Jurídica pelo titular", "Recebidos de Pessoa Jurídica pelos dependentes",
    "Recebidos de Pessoa Física/Exterior pelo titular", "Recebidos de Pessoa Física/Exterior pelos dependentes",
    "Recebidos acumuladamente pelo titular", "Recebidos acumuladamente pelos dependentes", "Resultado tributável da Atividade Rural",
]
PREPAID_LABELS = [
    ("t1", "Imposto retido na fonte do titular"), ("t2", "Imposto retido na fonte dos dependentes"),
    ("t3", "Carnê-Leão do titular"), ("t4", "Carnê-Leão dos dependentes"), ("t5", "Imposto complementar"),
    ("t6", "Imposto pago no exterior"), ("t7", "Imposto retido RRA"), ("t8", "Imposto retido na fonte (Lei nº 11.033/2004)"),
]

FONT, BLUE, DARK, TEAL = "Arial", "0000FF", "143E3C", "238783"
MONEY = 'R$ #,##0.00;(R$ #,##0.00);"-"'
PCT = '0.00%;(0.00%);"-"'
FLAG = '"Sim";"Sim";"Não"'
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="EDEDED")
SECTION_FILL = PatternFill("solid", fgColor=DARK)
SUB_FILL = PatternFill("solid", fgColor="DCEBEA")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
RESULT_FILL = PatternFill("solid", fgColor="FCE4D6")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
LAST_COL = 14  # A..N
WIDTHS = {"A": 7, "B": 44, "C": 17, "D": 17, "E": 17, "F": 17, "G": 17, "H": 11, "I": 11, "J": 11, "K": 16, "L": 10, "M": 15, "N": 46}


def consolidate_sources(raw):
    """Mesma regra de src/utils/dividends.ts: soma por CNPJ e agrupa fontes < R$ 500."""
    by_key, order = {}, []
    for s in raw:
        key = (s.get("cnpj") or s["name"], bool(s.get("simples")))
        if key in by_key:
            by_key[key]["total"] += s["total"]
            by_key[key]["protected"] += s.get("protected", 0)
        else:
            by_key[key] = {"name": s["name"], "cnpj": s.get("cnpj", ""), "total": s["total"], "protected": s.get("protected", 0),
                           "pj_profit": s.get("pj_profit", 0), "pj_tax": s.get("pj_tax", 0), "nominal": s.get("nominal", DEFAULT_NOMINAL),
                           "withholding": s.get("withholding", 0), "simples": bool(s.get("simples")), "note": s.get("note")}
            order.append(key)
    big = [by_key[k] for k in order if by_key[k]["total"] >= SMALL_SOURCE_THRESHOLD]
    small = [by_key[k] for k in order if by_key[k]["total"] < SMALL_SOURCE_THRESHOLD]
    big.sort(key=lambda s: -s["total"])
    for s in big:
        s["aggregate"] = False
        notes = [DEFAULT_ALLOCATION_NOTE if not s["protected"] else ""]
        if s["simples"]:
            notes.append("Distribuição de empresa do Simples Nacional (item 13) – também tem redutor.")
        if s["note"]:
            notes.append(s["note"])
        s["note"] = " ".join(n for n in notes if n)
    agg = None
    if small:
        agg = {"name": f"Demais fontes ({len(small)} fonte{'s' if len(small) > 1 else ''} abaixo de R$ {SMALL_SOURCE_THRESHOLD} cada)",
               "cnpj": "", "total": round(sum(s["total"] for s in small), 2), "protected": round(sum(s["protected"] for s in small), 2),
               "aggregate": True, "withholding": 0, "simples": False,
               "note": "Agrupadas: " + "; ".join(s["name"] for s in small) + ". Sem cálculo de redutor. " + DEFAULT_ALLOCATION_NOTE}
    return big, agg


def build_sheet(ws, p):
    ws.sheet_view.showGridLines = False
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width
    st = {"row": 1}
    R, LBL = {}, {}

    def style(cell, bold=False, color="000000", size=10, fill=None, fmt=None, align="left", wrap=False, italic=False):
        cell.font = Font(name=FONT, bold=bold, color=color, size=size, italic=italic)
        if fill:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)

    def merge(r, c1, c2):
        if c2 > c1:
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)

    def section(title):
        st["row"] += 1
        r = st["row"]
        merge(r, 1, LAST_COL)
        style(ws.cell(r, 1, title), bold=True, color="FFFFFF", size=11, fill=SECTION_FILL)
        ws.row_dimensions[r].height = 20
        st["row"] += 1

    def header(labels, spans=None):
        """labels: list of (col, text); spans: {col: last_col} for merged header cells."""
        r = st["row"]
        spans = spans or {}
        for col, text in labels:
            c = ws.cell(r, col, text)
            style(c, bold=True, color=DARK, fill=SUB_FILL, align="center" if col > 2 else "left", wrap=True)
            merge(r, col, spans.get(col, col))
        for col in range(1, LAST_COL + 1):
            ws.cell(r, col).border = BOX
            ws.cell(r, col).fill = SUB_FILL
        ws.row_dimensions[r].height = 30
        st["row"] += 1

    def finish(height=None):
        r = st["row"]
        for i in range(1, LAST_COL + 1):
            ws.cell(r, i).border = BOX
        if height:
            ws.row_dimensions[r].height = height
        st["row"] += 1

    A = lambda k: f"$C${R[k]}"

    # ---------------------------------------------------------------- title
    merge(1, 1, LAST_COL)
    ws["A1"] = "Simulação do IRPF Mínimo (Lei nº 15.270/2025) – execução da calculadora GSGA-PR"
    style(ws["A1"], bold=True, color=DARK, size=15)
    ws.row_dimensions[1].height = 24
    merge(2, 1, LAST_COL)
    ws["A2"] = (f"{p.get('confidential', 'CONFIDENCIAL')} – Contribuinte: {p['name']} – DIRPF 2026 (AC 2025). Uso interno GSGA-PR. "
                f"{p.get('header', '')} Fonte dos dados: {p.get('source', '')}")
    style(ws["A2"], italic=True, color="C00000", wrap=True)
    ws.row_dimensions[2].height = 34
    merge(3, 1, LAST_COL)
    ws["A3"] = ("Legenda: texto azul = dado de entrada · fundo amarelo = premissa ou dado a preencher/alocar · texto preto = fórmula. "
                "Cenário único (alocado): cada fonte pagadora separa a parcela deliberada até 2025 (protegida) da posterior (sujeita); "
                "o cálculo segue a alocação feita na tabela da seção 3.")
    style(ws["A3"], size=9, color="595959", wrap=True)
    ws.row_dimensions[3].height = 30
    st["row"] = 4

    # ---------------------------------------------------------------- 1. premissas
    section("1. PREMISSAS (alteráveis – fundo amarelo)")
    header([(1, "Ref."), (2, "Premissa"), (3, "Valor"), (4, "Fundamento / observação")], {4: LAST_COL})
    premises = [
        ("P1", "Limite inferior da base sujeita (alíquota zero até este valor)", 600000, MONEY,
         "Lei 15.270/2025 – rendimentos até R$ 600 mil no ano não sofrem o IRPF mínimo."),
        ("P2", "Limite superior da base sujeita (alíquota máxima a partir deste valor)", 1200000, MONEY,
         "Entre P1 e P2 a alíquota cresce linearmente de 0% até P3."),
        ("P3", "Alíquota máxima do IRPF mínimo", 0.10, PCT, "Alíquota aplicável a partir de R$ 1,2 milhão."),
        ("P4", "Alíquota presumida s/ rendimentos de tributação exclusiva/definitiva sem informação do imposto", 0.15, PCT,
         "Premissa do escritório: na falta de informação, presume-se 15%."),
        ("P5", "Alíquota nominal de IRPJ + CSLL padrão", DEFAULT_NOMINAL, PCT,
         "34% demais PJ · 40% seguradoras, capitalização e demais financeiras (CSLL 15%) · 45% bancos (CSLL 20%). "
         "Valor inicial da coluna H da tabela de fontes e base do indicador da linha C20."),
        ("P6", "Limite para agrupar fontes em \"Demais fontes\"", SMALL_SOURCE_THRESHOLD, MONEY,
         "Fontes que pagaram menos que isto cada uma são agrupadas numa única linha, sem redutor. Regra da ferramenta (informativa)."),
        ("P7", "Alíquota da retenção na fonte sobre dividendos posteriores a 2025", 0.10, PCT,
         "Aplicada à parcela posterior das fontes marcadas com retenção = Sim (coluna L). Padrão: Não."),
    ]
    extra = p.get("extra_subject")
    if extra:
        premises.append(("P8", extra["label"], extra["value"], MONEY, extra["note"]))
    for ref, label, value, fmt, why in premises:
        r = st["row"]
        R[ref] = r
        style(ws.cell(r, 1, ref), bold=True, color=DARK, align="center")
        style(ws.cell(r, 2, label), wrap=True)
        style(ws.cell(r, 3, value), color=BLUE, fill=YELLOW, fmt=fmt, align="center")
        merge(r, 4, LAST_COL)
        style(ws.cell(r, 4, why), size=9, color="595959", wrap=True)
        finish(30)

    qualitative = [
        ("Q1", "Base sujeita = rendimentos tributáveis + rendimentos de tributação exclusiva/definitiva + parcela POSTERIOR a 2025 dos lucros "
               "e dividendos" + (" + P8" if extra else "") + ". Demais isentos (poupança, LCI/LCA, CRI/CRA, FII, restituição de IR, doações e "
               "heranças declaradas) ficam fora da base."),
        ("Q2", "Alocação padrão: toda a distribuição de cada fonte é tratada como posterior a 2025 (hipótese conservadora). Havendo deliberação "
               "até 31/12/2025, informe a parcela protegida na coluna D da tabela de fontes; o cálculo se ajusta."),
        ("Q3", "Crédito contra o IRPF mínimo = antecipações + imposto sobre a tributação exclusiva (informado ou presumido) + retenção de 10% "
               "sobre dividendos (fontes marcadas). Não se usa o \"imposto devido na DAA\", para evitar dupla contagem."),
        ("Q4", "Redutor do art. 16-B, fonte a fonte = parcela posterior × máx(0; alíquota efetiva da PJ + alíquota do IRPF mínimo − alíquota "
               "nominal da PJ). Vale também para empresas do Simples Nacional. Soma limitada ao saldo devido; não gera restituição. "
               "Fonte sem dados da PJ ou agrupada em \"Demais fontes\" não gera redutor."),
        ("Q5", "Restituição: apenas a retenção na fonte sobre dividendos pode gerar restituição."),
        ("Q6", "A declaração é do AC 2025; a simulação aplica as regras de 2026 aos mesmos valores."),
    ]
    for ref, text in qualitative:
        r = st["row"]
        style(ws.cell(r, 1, ref), bold=True, color=DARK, align="center")
        merge(r, 2, LAST_COL)
        style(ws.cell(r, 2, text), wrap=True)
        finish(30)

    # ---------------------------------------------------------------- 2. dados
    section("2. DADOS DA DECLARAÇÃO (exceto lucros e dividendos – ver seção 3)")
    header([(1, "Ref."), (2, "Rendimento"), (3, "Valor (R$)"), (4, "Entra na base?"), (5, "Origem"), (6, "Observação")], {6: LAST_COL})
    counter = [0]

    def data_row(key, label, value, subject, origin, obs="", value_formula=False, fill=None):
        counter[0] += 1
        r = st["row"]
        LBL[key] = f"D{counter[0]}"
        R[key] = r
        style(ws.cell(r, 1, LBL[key]), bold=True, color=DARK, align="center")
        style(ws.cell(r, 2, label), wrap=True)
        style(ws.cell(r, 3, value), color="000000" if value_formula else BLUE, fmt=MONEY, align="right", fill=fill)
        if isinstance(subject, int):
            style(ws.cell(r, 4, subject), color=BLUE, fmt=FLAG, align="center")
        elif subject:
            style(ws.cell(r, 4, subject), size=9, align="center")
        style(ws.cell(r, 5, origin), size=9, wrap=True)
        merge(r, 6, LAST_COL)
        style(ws.cell(r, 6, obs), size=9, color="595959", wrap=True)
        finish(28 if len(obs) > 110 else None)

    trib_keys = []
    for i, lab in enumerate(TRIB_LABELS, 1):
        data_row(f"i{i}", lab, p.get("trib", {}).get(f"i{i}", 0), 1, "Resumo – Rend. tributáveis")
        trib_keys.append(f"i{i}")
    exempt_keys = []
    for i, (lab, v, origin, obs) in enumerate(p.get("exempt", [])):
        data_row(f"ex{i}", lab, v, 0, origin, obs)
        exempt_keys.append(f"ex{i}")
    excl_keys, foreign_key = [], None
    for i, e in enumerate(p.get("exclusive", [])):
        lab, v, origin, obs = e[:4]
        data_row(f"x{i}", lab, v, 1, origin, obs)
        excl_keys.append(f"x{i}")
        if len(e) > 4 and e[4]:
            foreign_key = f"x{i}"
    if extra:
        data_row("extra", f"{extra['label']} (premissa P8)", f"=$C${R['P8']}", 1, "Premissa do usuário (não consta da DIRPF)",
                 "Valor vem de P8.", value_formula=True)

    st["row"] += 1
    header([(1, "Ref."), (2, "Imposto pago / informado"), (3, "Valor (R$)"), (4, "Tratamento"), (5, "Origem"), (6, "Observação")], {6: LAST_COL})
    prep_keys = []
    for key, lab in PREPAID_LABELS:
        data_row(key, lab, p.get("prepaid", {}).get(key, 0), "Antecipação", "Resumo – Imposto pago")
        prep_keys.append(key)
    data_row("ftax", "Imposto sobre tributação exclusiva informado na DIRPF (ex.: Lei 14.754/2023)", p.get("foreign_tax", 0),
             "Trib. exclusiva", "Resumo – Imposto devido",
             f"Incide sobre {LBL[foreign_key]}; presunção P4 só sobre o restante." if foreign_key
             else "Não há nesta declaração → presunção P4 sobre toda a ficha exclusiva.")

    # ---------------------------------------------------------------- 3. fontes pagadoras
    big, agg = consolidate_sources(p.get("dividend_sources", []))
    section("3. LUCROS E DIVIDENDOS POR FONTE PAGADORA – alocação e redutor do art. 16-B (preencher colunas amarelas)")
    header([(1, "Ref."), (2, "Fonte pagadora (nome – CNPJ)"), (3, "Total recebido (R$)"), (4, "Deliberado até 2025 – protegido (R$)"),
            (5, "Posterior a 2025 – sujeito (R$)"), (6, "Lucro contábil da PJ antes de IRPJ/CSLL (R$)"), (7, "IRPJ + CSLL devidos (R$)"),
            (8, "Alíq. nominal"), (9, "Alíq. efetiva"), (10, "% redutor"), (11, "Redutor (R$)"), (12, "Retenção 10%?"),
            (13, "Retenção (R$)"), (14, "Observação")])
    nominal_dv = DataValidation(type="list", formula1='"0.34,0.4,0.45"', allow_blank=False)
    flag_dv = DataValidation(type="list", formula1='"0,1"', allow_blank=False)
    ws.add_data_validation(nominal_dv)
    ws.add_data_validation(flag_dv)
    first_src = st["row"]
    aliq_ref = "{ALIQ}"  # replaced after the calc rows exist
    src_rows = []

    def source_row(idx, s, aggregate=False):
        r = st["row"]
        src_rows.append(r)
        style(ws.cell(r, 1, f"F{idx}"), bold=True, color=DARK, align="center")
        name = s["name"] + (f" – {s['cnpj']}" if s.get("cnpj") else "") if s else ""
        style(ws.cell(r, 2, name), wrap=True, color=BLUE if s else "000000")
        style(ws.cell(r, 3, s["total"] if s else None), color=BLUE, fmt=MONEY, align="right")
        style(ws.cell(r, 4, s["protected"] if s else 0), color=BLUE, fill=YELLOW, fmt=MONEY, align="right")
        style(ws.cell(r, 5, f"=MAX(0,C{r}-D{r})"), fmt=MONEY, align="right")
        if aggregate:
            for c in (6, 7, 8, 9):
                style(ws.cell(r, c, "n/a"), color="7F7F7F", fill=GREY, align="center", size=9)
            style(ws.cell(r, 10, 0), fmt=PCT, align="right", fill=GREY)
        else:
            style(ws.cell(r, 6, s.get("pj_profit", 0) if s else 0), color=BLUE, fill=YELLOW, fmt=MONEY, align="right")
            style(ws.cell(r, 7, s.get("pj_tax", 0) if s else 0), color=BLUE, fill=YELLOW, fmt=MONEY, align="right")
            style(ws.cell(r, 8, s.get("nominal", DEFAULT_NOMINAL) if s else f"=$C${R['P5']}"), color=BLUE, fill=YELLOW, fmt=PCT, align="center")
            nominal_dv.add(f"H{r}")
            style(ws.cell(r, 9, f'=IF(F{r}>0,G{r}/F{r},"")'), fmt=PCT, align="right")
            style(ws.cell(r, 10, f"=IF(ISNUMBER(I{r}),MAX(0,I{r}+{aliq_ref}-H{r}),0)"), fmt=PCT, align="right")
        style(ws.cell(r, 11, f"=E{r}*J{r}"), fmt=MONEY, align="right")
        style(ws.cell(r, 12, s.get("withholding", 0) if s else 0), color=BLUE, fill=YELLOW, fmt=FLAG, align="center")
        flag_dv.add(f"L{r}")
        style(ws.cell(r, 13, f"=E{r}*L{r}*$C${R['P7']}"), fmt=MONEY, align="right")
        style(ws.cell(r, 14, (s or {}).get("note") or ("Linha livre para nova fonte." if not s else "")), size=9, color="595959", wrap=True)
        finish(42 if s and len(s.get("note") or "") > 120 else (30 if s else None))

    for i, s in enumerate(big, 1):
        source_row(i, s)
    for j in range(SPARE_SOURCE_ROWS):
        source_row(len(big) + j + 1, None)
    last_regular = st["row"] - 1
    if agg:
        source_row(len(big) + SPARE_SOURCE_ROWS + 1, agg, aggregate=True)
    last_src = st["row"] - 1

    r = st["row"]
    R["SRC_TOT"] = r
    style(ws.cell(r, 2, "TOTAL"), bold=True, fill=TOTAL_FILL)
    for col in (3, 4, 5, 11, 13):
        L = "ABCDEFGHIJKLMN"[col - 1]
        style(ws.cell(r, col, f"=SUM({L}{first_src}:{L}{last_src})"), bold=True, fmt=MONEY, align="right", fill=TOTAL_FILL)
    for col in (1, 6, 7, 8, 9, 10, 12, 14):
        ws.cell(r, col).fill = TOTAL_FILL
    finish()

    def check_row(key, label, formula, fmt, note, red_if_nonzero=True):
        r = st["row"]
        R[key] = r
        merge(r, 1, 2)
        style(ws.cell(r, 1, label), bold=True, size=9, wrap=True)
        style(ws.cell(r, 3, formula), fmt=fmt, align="right", bold=True,
              color=BLUE if isinstance(formula, (int, float)) else "000000")
        merge(r, 4, LAST_COL)
        style(ws.cell(r, 4, note), size=9, color="595959", wrap=True)
        if red_if_nonzero:
            ws.conditional_formatting.add(f"C{r}", CellIsRule(operator="notEqual", formula=["0"], fill=RED_FILL))
        finish()

    check_row("DECL", "Total de lucros e dividendos declarado na DIRPF (itens 09 + 13)", p.get("declared_dividend_total", 0), MONEY,
              "Dado de entrada (DIRPF). Usado para conferir a soma das fontes.", red_if_nonzero=False)
    check_row("CHK1", "Conferência: soma das fontes − total declarado", f"=ROUND(C{R['SRC_TOT']}-C{R['DECL']},2)", MONEY,
              "Deve ser zero (vermelho = diferença a investigar).")
    check_row("CHK2", "Conferência: fontes com parcela protegida maior que o total", f"=SUMPRODUCT((D{first_src}:D{last_src}>C{first_src}:C{last_src})*1)",
              '0', "Deve ser zero.")
    check_row("CHK3", "Fontes com dividendo sujeito sem dados da PJ (redutor zerado)",
              f"=SUMPRODUCT((E{first_src}:E{last_regular}>0)*(F{first_src}:F{last_regular}=0))", '0',
              "Preencher lucro e IRPJ + CSLL dessas PJs nas colunas F e G para calcular o redutor.")

    # ---------------------------------------------------------------- 4. cálculo
    section("4. CÁLCULO DO IRPF MÍNIMO (cenário alocado)")
    header([(1, "Ref."), (2, "Etapa"), (3, "Valor (R$)"), (4, "Memória de cálculo")], {4: LAST_COL})

    def calc(ref, label, formula, memo, fmt=MONEY, kind=None):
        r = st["row"]
        R[ref] = r
        fill = {"total": TOTAL_FILL, "result": RESULT_FILL}.get(kind)
        bold = kind in ("total", "result")
        style(ws.cell(r, 1, ref), bold=True, color=DARK, align="center", fill=fill)
        style(ws.cell(r, 2, label), bold=bold, fill=fill, wrap=True)
        style(ws.cell(r, 3, formula), bold=bold, fmt=fmt, align="right", fill=fill)
        merge(r, 4, LAST_COL)
        style(ws.cell(r, 4, memo), size=9, color="595959", wrap=True, fill=fill)
        finish()

    def sub(title):
        r = st["row"]
        merge(r, 1, LAST_COL)
        style(ws.cell(r, 1, title), bold=True, color=TEAL, fill=SUB_FILL)
        st["row"] += 1

    rng = lambda keys: (R[keys[0]], R[keys[-1]])
    lbl = lambda keys: f"{LBL[keys[0]]} a {LBL[keys[-1]]}" if len(keys) > 1 else LBL[keys[0]]
    tr, pr = rng(trib_keys), rng(prep_keys)
    T = R["SRC_TOT"]

    sub("4.1 Base sujeita ao IRPF mínimo")
    calc("C1", "Rendimentos tributáveis sujeitos", f"=SUMPRODUCT($C${tr[0]}:$C${tr[1]},$D${tr[0]}:$D${tr[1]})",
         f"Soma de {lbl(trib_keys)} marcados como \"Sim\".")
    if excl_keys:
        xr = rng(excl_keys)
        calc("C2", "(+) Rendimentos de tributação exclusiva/definitiva", f"=SUMPRODUCT($C${xr[0]}:$C${xr[1]},$D${xr[0]}:$D${xr[1]})",
             f"Soma de {lbl(excl_keys)} marcados como \"Sim\".")
    else:
        calc("C2", "(+) Rendimentos de tributação exclusiva/definitiva", 0, "Sem rendimentos de tributação exclusiva.")
    calc("C3", "(+) Lucros e dividendos posteriores a 2025", f"=E{T}", "Total da coluna E da tabela de fontes (seção 3).")
    base_parts = ["C1", "C2", "C3"]
    if extra:
        calc("C3a", f"(+) {extra['label']}", f"=$C${R['extra']}*$D${R['extra']}", f"{LBL['extra']} (premissa P8).")
        base_parts.append("C3a")
    calc("C4", "Base sujeita ao IRPF mínimo", "=" + "+".join(f"C{R[k]}" for k in base_parts), " + ".join(base_parts) + ".", kind="total")

    sub("4.2 Alíquota e IRPF mínimo devido")
    calc("C5", "Alíquota do IRPF mínimo",
         f"=IF(C{R['C4']}<={A('P1')},0,IF(C{R['C4']}>={A('P2')},{A('P3')},(C{R['C4']}-{A('P1')})/({A('P2')}-{A('P1')})*{A('P3')}))",
         "Se C4 ≤ P1: 0%. Se C4 ≥ P2: P3. Entre os dois: (C4 − P1) ÷ (P2 − P1) × P3.", fmt=PCT)
    calc("C6", "IRPF mínimo devido", f"=C{R['C4']}*C{R['C5']}", "C4 × C5.", kind="total")

    sub("4.3 Imposto já pago / considerado")
    calc("C7", "Antecipações (IRRF, carnê-leão, complementar etc.)", f"=SUM($C${pr[0]}:$C${pr[1]})", f"Soma de {lbl(prep_keys)}.")
    calc("C8", "Imposto sobre tributação exclusiva informado", f"={A('ftax')}", f"{LBL['ftax']}.")
    foreign = A(foreign_key) if foreign_key else "0"
    calc("C9", "Imposto presumido s/ tributação exclusiva sem informação",
         f"=MAX(0,C{R['C2']}-IF({A('ftax')}>0,{foreign},0))*{A('P4')}",
         "(C2 − rendimento cujo imposto foi informado) × P4.")
    calc("C10", "Retenção de 10% sobre dividendos posteriores (fontes marcadas)", f"=M{T}", "Total da coluna M da tabela de fontes.")
    calc("C11", "Total de imposto já pago / considerado", f"=SUM(C{R['C7']}:C{R['C10']})", "C7 + C8 + C9 + C10.", kind="total")
    calc("C12", "Saldo devido antes do redutor", f"=MAX(0,C{R['C6']}-C{R['C11']})", "máx(0; C6 − C11).")

    sub("4.4 Redutor do art. 16-B (fonte a fonte)")
    calc("C13", "Redutor potencial (soma das fontes)", f"=K{T}", "Total da coluna K da tabela de fontes.")
    calc("C14", "Redutor aplicado", f"=MIN(C{R['C12']},C{R['C13']})", "mín(C12; C13) – limitado ao saldo devido, sem restituição.",
         kind="total")

    sub("4.5 Resultado")
    calc("C15", "ADICIONAL DE IRPF MÍNIMO A PAGAR", f"=C{R['C12']}-C{R['C14']}", "C12 − C14.", kind="result")
    calc("C16", "Restituição", f"=MAX(0,C{R['C10']}-MAX(0,C{R['C6']}-(C{R['C11']}-C{R['C10']})))",
         "máx(0; C10 − máx(0; C6 − (C11 − C10))). Só a retenção sobre dividendos é restituível.", kind="result")
    calc("C17", "Carga total de IR", f"=C{R['C11']}+C{R['C15']}-C{R['C16']}", "C11 + C15 − C16.")
    calc("C18", "Carga efetiva sobre a base sujeita", f"=IF(C{R['C4']}>0,C{R['C17']}/C{R['C4']},0)", "C17 ÷ C4.", fmt=PCT)
    all_income = "+".join([f"SUM(C{R['i1']}:C{R['i7']})"] +
                          ([f"SUM(C{R[exempt_keys[0]]}:C{R[exempt_keys[-1]]})"] if exempt_keys else []) +
                          ([f"SUM(C{R[excl_keys[0]]}:C{R[excl_keys[-1]]})"] if excl_keys else []) +
                          ([f"C{R['extra']}"] if extra else []) + [f"C{T}"])
    calc("C19", "Carga efetiva sobre o rendimento total", f"=IF(({all_income})>0,C{R['C17']}/({all_income}),0)",
         "C17 ÷ (todos os rendimentos da seção 2 + total de dividendos).", fmt=PCT)
    calc("C20", "Alíquota efetiva uniforme das PJs que zeraria o adicional",
         f"=IF(AND(C{R['C3']}>0,C{R['C12']}>0),MIN(1,{A('P5')}-C{R['C5']}+C{R['C12']}/C{R['C3']}),0)",
         "P5 − C5 + C12 ÷ C3. Indicador: supõe todas as fontes com a mesma alíquota efetiva e nominal P5.", fmt=PCT)

    # resolve the IRPFM rate placeholder in the source table
    for r in src_rows:
        cell = ws.cell(r, 10)
        if isinstance(cell.value, str) and aliq_ref in cell.value:
            cell.value = cell.value.replace(aliq_ref, f"$C${R['C5']}")

    st["row"] += 1
    r = st["row"]
    merge(r, 1, LAST_COL)
    ws.cell(r, 1, "Simulação de caráter consultivo. Não substitui a declaração nem a análise individual do caso.")
    style(ws.cell(r, 1), size=8, italic=True, color="7F7F7F")

    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:3"
    return R


def build_workbook(people, out_path):
    wb = Workbook()
    for i, person in enumerate(people):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = person["sheet"][:31]
        build_sheet(ws, person)
    wb.calculation.fullCalcOnLoad = True
    wb.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as f:
        build_workbook(json.load(f), sys.argv[2])
    print("Planilha gerada:", sys.argv[2])
